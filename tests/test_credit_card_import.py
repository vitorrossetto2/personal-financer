import io
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

import app as novosaldo


def statement_bytes(include_new_purchase=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fatura"
    sheet["A1"] = "Vencimento"
    sheet["A2"] = "01/09/2026"
    sheet["B1"] = "Cartão"
    sheet["B2"] = "Cartão final 2788"
    sheet.append([])
    sheet.append(["Data", "Lan�amento", "Parcelamento", "Valor", "Nome", "N�mero do cart�o"])
    sheet.append(["03/08/2026", "Pagamento Efetuado", "", -200, "Vitor", "****2788"])
    sheet.append(["13/08/2026", "Mercado", "", 100, "Vitor", "****2788"])
    sheet.append(["10/08/2026", "Loja", "Parcela 2 de 5", 50, "Vitor", "****2788"])
    if include_new_purchase:
        sheet.append(["14/08/2026", "Café", "", 25, "Vitor", "****2788"])
    sheet.append([None, "Subtotal", None, "=SUBTOTAL(9,D6:D8)"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def future_statement_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Vencimento"
    sheet["A2"] = "01/10/2026"
    sheet["B1"] = "Cartão"
    sheet["B2"] = "Cartão final 2788"
    sheet.append([])
    sheet.append(["Data", "Lançamento", "Parcelamento", "Valor", "Nome", "Número do cartão"])
    sheet.append(["10/08/2026", "Loja", "Parcela 3 de 5", 55, "Vitor", "****2788"])
    sheet.append([None, "Subtotal", None, "=SUBTOTAL(9,D6:D6)"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parallel_installments_statement_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Vencimento"
    sheet["A2"] = "01/09/2026"
    sheet["B1"] = "Cartão"
    sheet["B2"] = "Cartão final 2788"
    sheet.append([])
    sheet.append(["Data", "Lançamento", "Parcelamento", "Valor", "Nome", "Número do cartão"])
    sheet.append(["10/08/2026", "Mesma loja", "Parcela 2 de 3", 80, "Vitor", "****2788"])
    sheet.append(["10/08/2026", "Mesma loja", "Parcela 2 de 3", 120, "Vitor", "****2788"])
    sheet.append([None, "Subtotal", None, "=SUBTOTAL(9,D6:D7)"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class CreditCardStatementTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        novosaldo.DATABASE = Path(self.directory.name) / "test.db"
        novosaldo.app.config.update(TESTING=True, SECRET_KEY="test", AUTH_DISABLED_FOR_TESTS=True)
        self.client = novosaldo.app.test_client()
        response = self.client.post(
            "/cartoes",
            data={"name": "Itaú Black", "issuer": "itau", "last_four": "2788", "payment_day": "5"},
        )
        self.assertEqual(response.status_code, 302)
        connection = sqlite3.connect(novosaldo.DATABASE)
        self.card_id = connection.execute("SELECT id FROM credit_cards").fetchone()[0]
        connection.close()

    def tearDown(self):
        self.directory.cleanup()

    def test_parser_uses_invoice_due_date_and_ignores_payment(self):
        parsed = novosaldo.parse_credit_card_statement(io.BytesIO(statement_bytes()), "itau")

        self.assertEqual(parsed["due_date"], "2026-09-01")
        self.assertEqual(parsed["skipped"], 1)
        self.assertEqual(len(parsed["entries"]), 2)
        self.assertEqual(sum(entry["amount"] for entry in parsed["entries"]), 150)
        self.assertEqual(parsed["entries"][1]["installment_number"], 2)
        self.assertEqual(parsed["entries"][1]["installment_total"], 5)
        self.assertEqual(parsed["entries"][1]["card_last_four"], "2788")

    def test_reimport_ignores_existing_rows_and_adds_only_new_purchases(self):
        first = self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(self.card_id), "billing_month": "2026-09",
                "category": "Cartão de crédito",
                "statement": (io.BytesIO(statement_bytes()), "fatura-aberta.xlsx"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        second = self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(self.card_id), "billing_month": "2026-09",
                "category": "Cartão de crédito",
                "statement": (io.BytesIO(statement_bytes(include_new_purchase=True)), "fatura-aberta.xlsx"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )

        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        connection = sqlite3.connect(novosaldo.DATABASE)
        current_rows = connection.execute(
            """SELECT description, amount, due_date, credit_card_id
               FROM transactions WHERE due_date BETWEEN '2026-09-01' AND '2026-09-30'
               ORDER BY amount DESC"""
        ).fetchall()
        installment_rows = connection.execute(
            """SELECT installment_number, due_date, card_installment_projected
               FROM transactions WHERE description = 'Loja' ORDER BY installment_number"""
        ).fetchall()
        connection.close()

        self.assertEqual(len(current_rows), 3)
        self.assertEqual({row[0] for row in current_rows}, {"Mercado", "Loja", "Café"})
        self.assertTrue(all(row[2] == "2026-09-05" for row in current_rows))
        self.assertTrue(all(row[3] == self.card_id for row in current_rows))
        self.assertEqual(
            installment_rows,
            [(2, "2026-09-05", 0), (3, "2026-10-05", 1),
             (4, "2026-11-05", 1), (5, "2026-12-05", 1)],
        )
        self.assertIn("2 já existia(m) neste cartão e mês", second.get_data(as_text=True))

    def test_reimport_backfills_projections_for_rows_with_legacy_import_keys(self):
        self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(self.card_id), "billing_month": "2026-09",
                "statement": (io.BytesIO(statement_bytes()), "fatura-legada.xlsx"),
            },
            content_type="multipart/form-data",
        )
        connection = sqlite3.connect(novosaldo.DATABASE)
        connection.execute("DELETE FROM transactions WHERE card_installment_projected = 1")
        connection.execute(
            "UPDATE transactions SET import_key = 'legacy-' || id WHERE credit_card_id = ?",
            (self.card_id,),
        )
        connection.commit()
        connection.close()

        response = self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(self.card_id), "billing_month": "2026-09",
                "statement": (io.BytesIO(statement_bytes()), "fatura-atual.xlsx"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )

        connection = sqlite3.connect(novosaldo.DATABASE)
        current_count = connection.execute(
            """SELECT COUNT(*) FROM transactions
               WHERE credit_card_id = ? AND due_date = '2026-09-05'""",
            (self.card_id,),
        ).fetchone()[0]
        installment_rows = connection.execute(
            """SELECT installment_number, due_date, card_installment_projected
               FROM transactions WHERE description = 'Loja' ORDER BY installment_number"""
        ).fetchall()
        connection.close()

        self.assertEqual(current_count, 2)
        self.assertEqual(
            installment_rows,
            [(2, "2026-09-05", 0), (3, "2026-10-05", 1),
             (4, "2026-11-05", 1), (5, "2026-12-05", 1)],
        )
        page = response.get_data(as_text=True)
        self.assertIn("2 já existia(m) neste cartão e mês", page)
        self.assertIn("3 parcela(s) futura(s) prevista(s)", page)

    def test_future_statement_confirms_projection_without_duplicate(self):
        self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(self.card_id), "billing_month": "2026-09",
                "statement": (io.BytesIO(statement_bytes()), "fatura-setembro.xlsx"),
            },
            content_type="multipart/form-data",
        )
        response = self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(self.card_id), "billing_month": "2026-10",
                "statement": (io.BytesIO(future_statement_bytes()), "fatura-outubro.xlsx"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )

        connection = sqlite3.connect(novosaldo.DATABASE)
        rows = connection.execute(
            """SELECT installment_number, amount, card_installment_projected, import_source
               FROM transactions WHERE description = 'Loja' ORDER BY installment_number"""
        ).fetchall()
        connection.close()

        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[1], (3, 55.0, 0, "fatura-outubro.xlsx"))
        self.assertIn("1 parcela(s) prevista(s) confirmada(s)", response.get_data(as_text=True))

    def test_parallel_installments_are_not_collapsed(self):
        self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(self.card_id), "billing_month": "2026-09",
                "statement": (
                    io.BytesIO(parallel_installments_statement_bytes()), "parcelas-paralelas.xlsx"
                ),
            },
            content_type="multipart/form-data",
        )
        connection = sqlite3.connect(novosaldo.DATABASE)
        rows = connection.execute(
            """SELECT installment_number, amount, card_installment_projected
               FROM transactions WHERE description = 'Mesma loja'
               ORDER BY installment_number, amount"""
        ).fetchall()
        connection.close()
        self.assertEqual(
            rows,
            [(2, 80.0, 0), (2, 120.0, 0), (3, 80.0, 1), (3, 120.0, 1)],
        )

    def test_manual_expense_uses_registered_payment_day(self):
        self.client.post(
            "/cartoes",
            data={"name": "Cartão dia 31", "issuer": "bradesco", "last_four": "9999", "payment_day": "31"},
        )
        connection = sqlite3.connect(novosaldo.DATABASE)
        card_id = connection.execute("SELECT id FROM credit_cards WHERE last_four = '9999'").fetchone()[0]
        connection.close()

        response = self.client.post(
            "/cartoes/lancamentos",
            data={
                "card_id": str(card_id), "billing_month": "2027-02", "description": "Compra manual",
                "category": "Teste", "amount": "30,00", "purchase_date": "2027-01-20",
            },
        )
        self.assertEqual(response.status_code, 302)
        connection = sqlite3.connect(novosaldo.DATABASE)
        due_date = connection.execute(
            "SELECT due_date FROM transactions WHERE description = 'Compra manual'"
        ).fetchone()[0]
        connection.close()
        self.assertEqual(due_date, "2027-02-28")

    def test_itau_statement_must_match_registered_card_final(self):
        self.client.post(
            "/cartoes",
            data={"name": "Outro Itaú", "issuer": "itau", "last_four": "1234", "payment_day": "1"},
        )
        connection = sqlite3.connect(novosaldo.DATABASE)
        other_card_id = connection.execute("SELECT id FROM credit_cards WHERE last_four = '1234'").fetchone()[0]
        connection.close()
        response = self.client.post(
            "/cartoes/importar",
            data={
                "card_id": str(other_card_id), "billing_month": "2026-09",
                "statement": (io.BytesIO(statement_bytes()), "fatura.xlsx"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertIn("não do final 1234", response.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
