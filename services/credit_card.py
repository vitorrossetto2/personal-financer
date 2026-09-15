from __future__ import annotations

import calendar
import hashlib
import re
import sqlite3
import unicodedata
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook

from repositories import CreditCardRepository, TransactionRepository


def normalize_label(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", text).strip().casefold()


def statement_header(value) -> str | None:
    """Mapeia também os cabeçalhos corrompidos presentes no XLSX exportado pelo Itaú."""
    label = re.sub(r"[^a-z0-9 ]", "", normalize_label(value))
    if label == "data":
        return "data"
    if label.startswith("lan") and label.endswith("amento"):
        return "lancamento"
    if label == "parcelamento":
        return "parcelamento"
    if label == "valor":
        return "valor"
    if label == "nome":
        return "nome"
    if "mero do cart" in label or label.startswith("nmero do cart"):
        return "numero do cartao"
    return None


def spreadsheet_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def spreadsheet_amount(value) -> Decimal | None:
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    text = str(value or "").strip().replace("R$", "").replace(" ", "")
    if not text or text.startswith("="):
        return None
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except Exception:
        return None


def shifted_month(billing_month: str, offset: int) -> str:
    first = datetime.strptime(billing_month, "%Y-%m").date()
    year = first.year + (first.month - 1 + offset) // 12
    month = (first.month - 1 + offset) % 12 + 1
    return f"{year:04d}-{month:02d}"


def entry_identity(entry: dict, billing_month: str, occurrence: int) -> str:
    """Identidade estável que faz a parcela prevista coincidir com a fatura futura."""
    installment_number = entry.get("installment_number")
    installment_total = entry.get("installment_total")
    amount_identity = "" if installment_number else str(
        Decimal(str(entry["amount"])).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    )
    parts = (
        billing_month,
        entry["purchase_date"],
        normalize_label(entry["description"]),
        amount_identity,
        f"{installment_number or ''}/{installment_total or ''}",
        entry.get("card_last_four") or "",
        normalize_label(entry.get("cardholder_name")),
        str(occurrence),
    )
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()


def installment_group_identity(entry: dict, occurrence: int) -> str:
    parts = (
        entry["purchase_date"], normalize_label(entry["description"]),
        entry.get("card_last_four") or "", normalize_label(entry.get("cardholder_name")),
        str(entry.get("installment_total") or ""), str(occurrence),
    )
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()[:24]


def parse_itau_statement(file_object) -> dict:
    """Lê a planilha de fatura do Itaú sem interpretar seu texto como instruções."""

    def row_value(row, column):
        return row[column] if column is not None and column < len(row) else None

    workbook = load_workbook(file_object, read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            header_row = None
            columns = {}
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                labels = {
                    key: index for index, value in enumerate(row)
                    if value is not None and (key := statement_header(value))
                }
                if {"data", "lancamento", "valor"}.issubset(labels):
                    header_row = row_number
                    columns = labels
                    break
            if not header_row:
                continue

            due_date = None
            card_label = None
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=header_row - 1, values_only=True), start=1
            ):
                for column_number, value in enumerate(row, start=1):
                    compact_label = re.sub(r"[^a-z0-9]", "", normalize_label(value))
                    if compact_label == "vencimento":
                        for candidate_row in range(row_number + 1, min(header_row, row_number + 5)):
                            due_date = spreadsheet_date(sheet.cell(candidate_row, column_number).value)
                            if due_date:
                                break
                    elif compact_label in {"cartao", "carto"}:
                        for candidate_row in range(row_number + 1, min(header_row, row_number + 5)):
                            candidate = sheet.cell(candidate_row, column_number).value
                            if candidate:
                                card_label = str(candidate).strip()
                                break

            if not due_date:
                raise ValueError("Não foi possível identificar o vencimento da fatura do Itaú.")

            entries = []
            skipped = 0
            occurrences = Counter()
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                description = str(row_value(row, columns["lancamento"]) or "").strip()
                if normalize_label(description).startswith("subtotal") or any(
                    normalize_label(value).startswith("subtotal") for value in row if value is not None
                ):
                    break
                purchase_date = spreadsheet_date(row_value(row, columns["data"]))
                amount = spreadsheet_amount(row_value(row, columns["valor"]))
                if not description and amount is None:
                    continue
                if not purchase_date or amount is None or amount <= 0:
                    skipped += 1
                    continue

                installment_text = str(row_value(row, columns.get("parcelamento")) or "").strip()
                match = re.search(r"(?:parcela\s*)?(\d+)\s*(?:de|/)\s*(\d+)", installment_text, re.I)
                card_number = str(row_value(row, columns.get("numero do cartao")) or "").strip()
                owner = str(row_value(row, columns.get("nome")) or "").strip()
                installment_number = int(match.group(1)) if match else None
                installment_total = int(match.group(2)) if match else None
                card_last_four = re.sub(r"\D", "", card_number)[-4:] or None
                entry = {
                    "description": description,
                    "amount": float(amount),
                    "purchase_date": purchase_date.isoformat(),
                    "installment_number": installment_number,
                    "installment_total": installment_total,
                    "cardholder_name": owner or None,
                    "card_last_four": card_last_four,
                }
                occurrence_key = (
                    entry["purchase_date"], normalize_label(description),
                    "" if installment_number else str(amount),
                    installment_number, installment_total, card_last_four, normalize_label(owner),
                )
                occurrences[occurrence_key] += 1
                occurrence = occurrences[occurrence_key]
                entry["occurrence"] = occurrence
                entry["installment_group"] = (
                    installment_group_identity(entry, occurrence) if match else None
                )
                entry["import_key"] = entry_identity(entry, due_date.strftime("%Y-%m"), occurrence)
                entries.append(entry)
            if not entries:
                raise ValueError("A planilha não contém compras válidas para importar.")
            return {
                "due_date": due_date.isoformat(), "card_label": card_label,
                "entries": entries, "skipped": skipped,
            }
    finally:
        workbook.close()
    raise ValueError("Não encontrei as colunas Data, Lançamento e Valor na planilha do Itaú.")


def parse_credit_card_statement(file_object, issuer: str = "itau") -> dict:
    parser = {"itau": parse_itau_statement}.get(issuer)
    if not parser:
        raise ValueError(f"A importação de faturas do emissor {issuer.title()} ainda não está disponível.")
    return parser(file_object)


class CreditCardService:
    def __init__(self, connection: sqlite3.Connection):
        self.connection = connection
        self.cards = CreditCardRepository(connection)
        self.transactions = TransactionRepository(connection)

    def create_card(self, *, name: str, issuer: str, last_four: str, payment_day: int) -> int:
        name = name.strip()
        issuer = issuer.strip().casefold()
        last_four = re.sub(r"\D", "", last_four)
        if not name or issuer not in {"itau", "bradesco"}:
            raise ValueError("Informe o nome e um emissor válido.")
        if len(last_four) != 4:
            raise ValueError("Informe os quatro últimos dígitos do cartão.")
        if payment_day < 1 or payment_day > 31:
            raise ValueError("O dia de pagamento deve estar entre 1 e 31.")
        timestamp = datetime.now().isoformat(timespec="seconds")
        card_id = self.cards.create(
            name=name, issuer=issuer, last_four=last_four,
            payment_day=payment_day, timestamp=timestamp,
        )
        self.connection.commit()
        return card_id

    def list_cards(self):
        return self.cards.list_all()

    @staticmethod
    def payment_date(card, billing_month: str) -> date:
        try:
            first = datetime.strptime(billing_month, "%Y-%m").date().replace(day=1)
        except ValueError as error:
            raise ValueError("Informe o mês da fatura no formato AAAA-MM.") from error
        day = min(card["payment_day"], calendar.monthrange(first.year, first.month)[1])
        return first.replace(day=day)

    def create_expense(self, *, card_id: int, billing_month: str, description: str,
                       category: str, amount: float, purchase_date: date) -> int:
        card = self.cards.find_active(card_id)
        if not card:
            raise ValueError("Cartão não encontrado.")
        description = description.strip()
        if not description or amount <= 0:
            raise ValueError("Informe descrição e valor válidos.")
        transaction_id = self.transactions.create(
            kind="expense", description=description, category=category.strip(), amount=amount,
            due_date=self.payment_date(card, billing_month).isoformat(),
            purchase_date=purchase_date.isoformat(), credit_card_id=card_id,
            created_at=datetime.now().isoformat(timespec="seconds"),
        )
        self.connection.commit()
        return transaction_id

    def control_page(self, card_id: int | None, billing_month: str, today: date) -> dict:
        cards = self.cards.list_all()
        selected_card = self.cards.find_active(card_id) if card_id else (cards[0] if cards else None)
        if card_id and not selected_card:
            raise ValueError("Cartão não encontrado.")
        first = datetime.strptime(billing_month, "%Y-%m").date().replace(day=1)
        last = first.replace(day=calendar.monthrange(first.year, first.month)[1])
        rows = [] if not selected_card else self.transactions.list_for_credit_card_period(
            selected_card["id"], first.isoformat(), last.isoformat(), today.isoformat()
        )
        return {
            "cards": cards,
            "selected_card": selected_card,
            "billing_month": billing_month,
            "rows": rows,
            "total": round(sum(row["amount"] for row in rows), 2),
            "payment_date": self.payment_date(selected_card, billing_month).isoformat()
            if selected_card else None,
        }

    def save_statement(self, statement: dict, filename: str, *, credit_card_id: int,
                       category: str = "Cartão de crédito") -> dict:
        card = self.cards.find_active(credit_card_id)
        if not card:
            raise ValueError("Cartão não encontrado.")
        statement_digits = re.sub(r"\D", "", statement.get("card_label") or "")
        if statement_digits and not statement_digits.endswith(card["last_four"]):
            raise ValueError(
                f"A fatura é do cartão final {statement_digits[-4:]}, não do final {card['last_four']}."
            )
        category = category.strip() or "Cartão de crédito"
        created_at = datetime.now().isoformat(timespec="seconds")
        billing_month = statement["due_date"][:7]
        target_due_date = self.payment_date(card, billing_month).isoformat()
        imported = 0
        confirmed = 0
        duplicates = 0
        scheduled = 0
        for original_entry in statement["entries"]:
            entry = dict(original_entry)
            occurrence = int(entry.get("occurrence") or 1)
            entry["import_key"] = hashlib.sha256(
                f"{credit_card_id}|{entry_identity(entry, billing_month, occurrence)}".encode("utf-8")
            ).hexdigest()
            result = self.transactions.import_card_entry(
                entry, credit_card_id=credit_card_id, due_date=target_due_date,
                category=category, filename=filename, created_at=created_at,
            )
            if result == "inserted":
                imported += 1
            elif result == "confirmed":
                confirmed += 1
            else:
                duplicates += 1

            installment_number = entry.get("installment_number")
            installment_total = entry.get("installment_total")
            if not installment_number or not installment_total or installment_number >= installment_total:
                continue
            for offset in range(1, installment_total - installment_number + 1):
                future_entry = dict(entry)
                future_entry["installment_number"] = installment_number + offset
                future_month = shifted_month(billing_month, offset)
                future_entry["import_key"] = hashlib.sha256(
                    f"{credit_card_id}|{entry_identity(future_entry, future_month, occurrence)}".encode("utf-8")
                ).hexdigest()
                future_result = self.transactions.import_card_entry(
                    future_entry, credit_card_id=credit_card_id,
                    due_date=self.payment_date(card, future_month).isoformat(),
                    category=category, filename=filename, created_at=created_at, projected=True,
                )
                if future_result == "inserted":
                    scheduled += 1
        self.connection.commit()
        return {
            "imported": imported,
            "confirmed": confirmed,
            "duplicates": duplicates,
            "scheduled": scheduled,
            "skipped": statement["skipped"],
            "due_date": target_due_date,
            "total": round(sum(entry["amount"] for entry in statement["entries"]), 2),
        }
